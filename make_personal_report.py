# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\\unity_project\\CS\\My project")
TEMPLATE = ROOT / "网络工程项目实施-个人报告模板.docx"
OUTPUT = ROOT / "网络工程项目实施-个人报告_服务端方向.docx"


def clear_para(p, value):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)
    p.add_run(value)


def add_after(doc, target, value, style="Normal"):
    p = doc.add_paragraph(value, style=style)
    target._p.addnext(p._p)
    return p


def set_font(run, name="宋体", size=12, bold=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), name)


def style_font(style, size, bold=False):
    style.font.name = "宋体"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), "宋体")


def format_body(p):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(5)


def fill_slot(doc, body_para, texts):
    clear_para(body_para, texts[0])
    current = body_para
    for value in texts[1:]:
        current = add_after(doc, current, value)
    return current


doc = Document(TEMPLATE)
section = doc.sections[0]
section.top_margin = Inches(0.5)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.5)
section.right_margin = Inches(0.5)

# 统一模板字体，保留模板的标题层级和页面结构。
style_font(doc.styles["Title"], 20, True)
style_font(doc.styles["Subtitle"], 16, True)
style_font(doc.styles["副标题2"], 12, False)
style_font(doc.styles["Normal"], 12, False)
style_font(doc.styles["项目"], 14, True)

paras = list(doc.paragraphs)
clear_para(paras[0], "《网络工程项目实施》个人报告")
clear_para(paras[1], "Unity 多人合作射击项目")
clear_para(paras[2], "第x组 组长（服务端方向）")
clear_para(paras[3], "本报告以我在 Unity 多人合作射击项目中的实际工作为主线，说明我负责的服务端网络模块、完成过程、能力提升和后续改进方向。")

role = [
    "在本项目中，我主要承担服务端网络接入、房间流程、权威状态、战斗命令和快照广播等工作。多人模式的开局、战斗和结束都需要这些模块衔接，因此我也承担了部分组长协调工作，负责把客户端、怪物和界面需要的状态整理成统一的消息和接口。",
    "我的工作重点不是单独完成某一个界面，而是保证服务器能够稳定接收客户端请求，按照固定的游戏节奏推进状态，并把可验证的结果返回给各个客户端。这样可以减少不同客户端各自计算造成的血量、弹药和击杀分数不一致。",
]

contribution = [
    "第一，我实现了 Python asyncio 服务端的 TCP 和 UDP 接入。TCP 使用换行分隔的 JSON 消息，负责大厅、hello/welcome、ready/start、地图上传、射击、换弹和击杀事件；UDP 负责移动输入、绑定和实时状态快照。接收网络数据的部分只负责解析和排队，游戏循环统一处理状态，避免多个网络任务同时修改玩家和怪物对象。",
    "第二，我整理了多人模式的房间流程。客户端先发送 hello，服务器返回玩家编号和大厅状态；房主提交 ready/start 后，服务器接收并检查地图数据，生成第一波怪物，广播 game_started，最后绑定 UDP 进入实时游戏。断开连接时，服务器会清理玩家、远端对象和本局状态，避免下一次进入游戏时残留旧数据。",
    "第三，我负责维护服务器的权威状态。服务器保存玩家位置、输入序号、生命代次、生命值、弹药、动作计时和分数，也保存怪物位置、血量、速度、攻击冷却和击退状态。服务器只接受新的序号，拒绝旧的移动、射击和换弹命令；玩家死亡后还会增加 life，防止上一条生命的命令在复活后继续生效。",
    "第四，我完成了战斗命令和快照广播。服务器收到射击命令后检查射速、弹药、方向、生命代次和地图遮挡，命中后扣除怪物血量并施加小幅击退；怪物死亡时增加击杀者 10 分，并通过事件和快照通知客户端。怪物近身攻击按每秒 10 点扣除玩家生命，玩家在 2 秒倒计时后由服务器选择安全位置满血复活，并恢复初始弹药。",
    "第五，我将游戏循环控制在约 30Hz，状态快照控制在约 20Hz。快照包含玩家、怪物、生命、弹药、波次和排行榜等信息，客户端只负责表现。针对 UDP 可能出现的乱序和重复，我在客户端配合使用 tick 和分数序号过滤旧状态，减少弹药回退、排行榜倒退和怪物重复出现等问题。",
]

improve = [
    "通过这次开发，我对 TCP 和 UDP 的使用场景有了更具体的认识。大厅、开局、地图、射击和击杀事件需要可靠到达，适合使用 TCP；移动和高频快照允许丢弃过期数据，适合使用 UDP。换行分隔 JSON 也让我理解了粘包、半包和消息边界在实际程序中的影响。",
    "我对 asyncio 的任务协作、队列和固定频率游戏循环更加熟悉。以前我更容易把网络接收和游戏逻辑写在一起，这次通过“接收、排队、统一模拟、广播”的结构，能够更清楚地区分数据入口和状态修改位置。",
    "我还提高了排查联机问题的能力。弹药变少又变多、死亡后旧命令生效、怪物状态回退等问题，不能只看画面，需要同时检查 tick、seq、life 和服务器日志。通过和客户端同学联调，我逐渐形成了先确认消息顺序，再确认状态归属，最后检查表现层的排查方法。",
]

need_more = [
    "目前服务端主要面向同机或局域网运行，断线重连、房间持久化和公网部署还不够完善。后续我需要补充连接恢复流程、房间编号和版本协商，并增加对异常客户端退出的压力测试。",
    "当怪物数量增加时，寻路、碰撞检查和 JSON 快照都会增加服务器负担。后续可以统计每帧寻路和序列化耗时，减少重复计算，并考虑更紧凑的状态格式。对地图和怪物参数也应增加自动化测试，避免客户端采集数据异常时影响整局游戏。",
    "在个人能力方面，我还需要加强网络安全和工程化方面的知识，例如更严格的输入校验、日志分级、限流和错误恢复。现在的实现能够支撑课程项目演示，但距离可长期运行的服务还需要更多测试和监控。",
]

ai_use = [
    "开发过程中，我使用 AI 辅助梳理网络流程、检查 Python 和 C# 代码结构，并根据报错信息寻找可能的原因。它在整理 TCP/UDP 分工、状态字段和报告表达方面比较高效，也能帮助我快速比较不同实现方案。",
    "但 AI 不了解项目运行时的全部状态，生成的代码可能和现有脚本、Unity 生命周期或消息字段不匹配。比如网络状态回退、对象重复生成这类问题，不能只凭一段代码判断，必须结合实际日志、客户端画面和服务端行为验证。因此我把 AI 当作辅助工具，最终的接口设计、代码合并和运行测试由我和组员共同确认。",
]

team = [
    "作为服务端方向的负责人，我先把 hello、ready/start、map、shoot、ammo_action、snapshot 和 event 等消息类型整理出来，明确字段含义、发送方向和状态更新时机，再和客户端、怪物及 HUD 模块的同学对接。这样可以减少每个人按照自己的理解修改消息造成的反复返工。",
    "团队协作中，我把问题分成网络接入、玩法状态、界面表现和场景碰撞几类，联调时优先处理会阻断整局游戏的问题，例如无法进入房间、快照不更新和玩家状态不同步，再处理特效、排行榜和提示文字等表现问题。每次完成一项功能后，我们都会用单机和多人两种模式分别验证，确认修改没有破坏已有功能。",
    "我认为组长的作用不只是分配任务，还要保持接口和进度透明。后续如果继续开发，我会把消息协议、测试步骤和已知问题整理成简短文档，让组员可以快速定位自己的改动范围，提高多人协作效率。",
]

fill_slot(doc, paras[6], role)
fill_slot(doc, paras[8], contribution)
fill_slot(doc, paras[10], improve)
fill_slot(doc, paras[12], need_more)
fill_slot(doc, paras[14], ai_use)

# 模板最后一个标题已经是团队协作问题，正文直接接在标题后面。
last = doc.paragraphs[-1]
current = add_after(doc, last, team[0])
for value in team[1:]:
    current = add_after(doc, current, value)

for p in doc.paragraphs:
    if p.style.name == "项目":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    elif p.style.name == "Normal":
        format_body(p)
    elif p.style.name in ("Title", "Subtitle", "副标题2"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
    for run in p.runs:
        if p.style.name == "Title":
            set_font(run, size=20, bold=True)
        elif p.style.name == "Subtitle":
            set_font(run, size=16, bold=True)
        elif p.style.name == "项目":
            set_font(run, size=14, bold=True)
        else:
            set_font(run, size=12, bold=False)

doc.core_properties.title = "Unity 多人合作射击项目个人报告"
doc.core_properties.subject = "网络工程项目实施课程项目个人总结"
doc.core_properties.author = "第x组组长（服务端方向）"
doc.save(OUTPUT)
print(OUTPUT)
