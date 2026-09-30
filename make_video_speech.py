# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

ROOT = Path(r"D:\\unity_project\\CS\\My project")
OUT = ROOT / "Unity多人合作射击项目_视频演讲稿.docx"


def set_font(run, size=12, bold=False, italic=False):
    run.font.name = "宋体"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), "宋体")


def set_style(style, size, bold=False):
    style.font.name = "宋体"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), "宋体")


def add_para(doc, text, cue=False):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if cue else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, 11 if cue else 12, italic=cue)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_font(run, 14, bold=True)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
set_style(doc.styles["Normal"], 12)
set_style(doc.styles["Title"], 20, True)
set_style(doc.styles["Subtitle"], 14)
set_style(doc.styles["Heading 1"], 14, True)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
set_font(title.add_run("Unity 多人合作射击项目视频演讲稿"), 20, bold=True)

subtitle = doc.add_paragraph(style="Subtitle")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(18)
set_font(subtitle.add_run("项目内容、运行效果与设计亮点"), 14)

add_para(doc, "使用说明：正文可以直接照着录制，方括号中的内容是画面操作提示，不需要念出来。整段演示控制在 6 到 8 分钟较合适。", True)

add_heading(doc, "一、开场介绍")
add_para(doc, "大家好，我们小组这次课程项目的名称是“Unity 多人合作射击项目”。")
add_para(doc, "这是一个基于 Unity 开发的第一人称合作射击游戏。玩家可以选择单机模式，也可以连接本地服务器进入多人模式，和其他玩家一起对抗不断出现的怪物。")
add_para(doc, "项目的重点不只是实现射击功能，还包括多人网络通信、服务器权威状态、客户端平滑处理、怪物寻路、玩家生命与复活、弹药系统以及 HUD 界面等内容。接下来我会从项目背景、整体架构、功能演示和设计亮点几个方面进行介绍。")

add_heading(doc, "二、项目整体设计")
add_para(doc, "[展示项目主菜单或系统架构图]", True)
add_para(doc, "我们的项目主要分为 Unity 客户端和 Python 服务端两部分。Unity 客户端负责场景显示、玩家操作、武器表现、怪物模型、界面和特效；多人模式下，Python 服务端负责保存和计算真正的游戏状态，包括玩家位置、生命值、弹药、怪物血量、怪物攻击、击杀分数和波次进度等。")
add_para(doc, "通信方面，我们将 TCP 和 UDP 分开使用。TCP 负责必须可靠到达的消息，例如进入房间、开始游戏、地图上传、射击、换弹和击杀事件；UDP 负责移动输入和实时状态快照，因为这些数据更新频率较高，过期的数据可以直接丢弃。当前服务端主要运行在本机或局域网中，其他电脑可以通过服务器的局域网 IP 地址连接。")

add_heading(doc, "三、单机模式演示")
add_para(doc, "[进入单机模式，设置怪物数量和初始血量]", True)
add_para(doc, "首先演示单机模式。在开始界面中，玩家可以设置怪物数量和初始血量。点击开始后，玩家会进入战斗场景。玩家使用 WASD 移动，鼠标控制视角，鼠标左键进行射击。当前版本支持连续射击，按住鼠标左键可以按照武器射速持续开火。")
add_para(doc, "[移动并射击怪物]", True)
add_para(doc, "怪物受到子弹攻击后会减少生命值，同时受到小幅度击退。当怪物生命值归零后，会播放死亡爆炸特效，然后从场景中消失。怪物不会一直沿直线移动，而是根据地图碰撞网格和障碍物进行方向判断，在墙边或障碍物附近尝试绕行。如果怪物长时间无法移动，系统会重新寻找可行位置，减少卡在墙角的情况。")

add_heading(doc, "四、生命、死亡和复活功能")
add_para(doc, "[让角色受到攻击，展示左下角血条]", True)
add_para(doc, "屏幕左下角是玩家血条，血量下降时，血条的长度会按照剩余生命值比例缩短，同时显示当前血量。怪物靠近玩家后，会在攻击范围内造成伤害，当前设定是每秒受到 10 点伤害。")
add_para(doc, "当玩家生命值降到 0 后，角色会进入死亡状态。死亡期间画面固定在死亡位置，镜头逐渐向下，模拟倒地效果，同时显示 2 秒复活倒计时和进度条。复活时，系统会选择一个安全的随机位置，玩家恢复满血并重新获得初始弹药。")

add_heading(doc, "五、弹药和换弹系统演示")
add_para(doc, "[展示弹药 HUD，连续射击并进行换弹]", True)
add_para(doc, "屏幕右下角显示当前弹夹和备用弹药数量，初始化为 50 发当前弹夹和 200 发备用弹药。当当前弹夹打空并且还有备用弹药时，会自动进入换弹流程，换弹时间大约为 1.5 秒；玩家也可以短按一次 R 键主动换弹。")
add_para(doc, "如果需要恢复完整弹药，可以长按 R 键 2 秒。这个操作会同时补充当前弹夹和备用弹药，进度条会随着按键持续时间变化，达到 2 秒后完成补充。这样换弹和补充弹药是两种不同的操作，玩家可以根据战斗情况选择合适的方式。")

add_heading(doc, "六、三波怪物和 HUD 功能")
add_para(doc, "[展示怪物波次提示]", True)
add_para(doc, "本项目将简单的一波出怪扩展成了三波怪物。每波之间间隔 20 秒，在下一波开始前 5 秒，屏幕上会显示倒计时和即将到来的波次。后续波次相比前一波会在数量、血量、伤害和移动速度上进行小幅提升，使战斗难度逐渐增加。")
add_para(doc, "[展示小地图]", True)
add_para(doc, "屏幕左上角是圆形小地图：绿色三角形表示自己，蓝色三角形表示队友，红色点表示怪物，三角形方向会跟随角色当前朝向变化。小地图可以帮助玩家快速判断队友和怪物的位置。")
add_para(doc, "[展示排行榜]", True)
add_para(doc, "多人模式左侧中部有击杀排行榜。每击杀一只怪物获得 10 分，排行榜根据玩家分数排序。分数由服务器统一计算并广播，客户端只负责显示，避免不同玩家看到的分数不一致。")

add_heading(doc, "七、多人模式演示")
add_para(doc, "[启动本地服务器，打开两个客户端]", True)
add_para(doc, "进入多人模式后，玩家需要输入服务器地址和昵称。连接成功后，客户端先通过 TCP 发送 hello 消息，服务器返回玩家编号和大厅状态。房主点击开始后，客户端将地图碰撞信息发送给服务器，服务器验证地图数据、生成第一波怪物，并通知所有客户端进入游戏。")
add_para(doc, "[展示两个客户端同步移动和射击]", True)
add_para(doc, "在游戏过程中，服务器负责计算最终状态。一个客户端移动或射击后，服务器会处理对应命令，再通过快照和事件把结果发送给所有客户端。这样玩家的生命值、弹药、怪物血量、击杀分数和波次状态可以在多个客户端之间保持一致。")

add_heading(doc, "八、网络平滑处理设计")
add_para(doc, "多人游戏中，如果角色完全等待服务器返回位置，操作会有明显延迟。因此我们在客户端加入了平滑处理。本地玩家移动时，客户端会先根据输入进行预测，不需要每次等待服务器返回结果；客户端和服务器都按照约 30Hz 的频率推进移动。")
add_para(doc, "服务器通过 UDP 返回带有 tick 和序号的状态快照。客户端收到后，会将本地预测位置和服务器权威位置进行比较：误差较小时保持连续移动，误差中等时逐渐校正，误差过大或玩家刚刚复活时直接接受服务器位置。")
add_para(doc, "对于其他玩家和怪物，客户端不会收到快照后直接瞬移，而是使用插值、平滑移动和旋转插值，让远端对象的动作更加自然。这种处理方式在保证服务器最终裁决的同时，减少了网络延迟带来的卡顿和瞬移。")

add_heading(doc, "九、项目设计亮点")
add_para(doc, "第一是服务器权威状态。玩家的生命、弹药、命中、怪物血量和击杀分数都由服务器最终决定，客户端不能直接修改这些结果，减少了状态不同步问题。")
add_para(doc, "第二是 TCP 和 UDP 的分工。我们让 TCP 负责可靠事件，让 UDP 负责高频实时状态，从而同时满足可靠性和实时性的要求。")
add_para(doc, "第三是客户端平滑处理。客户端对本地玩家进行预测，对远端玩家和怪物进行插值，并按照误差大小进行分段校正，降低网络延迟对操作体验的影响。")
add_para(doc, "第四是单机和多人共用玩法逻辑。单机模式使用 Unity 本地逻辑，多人模式使用服务器逻辑，但两种模式都保留了相同的生命、弹药、怪物、波次和战斗规则。")
add_para(doc, "第五是功能比较完整。项目不仅有基础射击，还加入了怪物绕障、击退、死亡爆炸、三波出怪、生命和复活、弹药补充、小地图、击杀排行榜以及灵敏度设置等功能。")

add_heading(doc, "十、项目总结")
add_para(doc, "通过这次项目，我们完成了从单机射击原型到局域网多人合作游戏的扩展。在开发过程中，我们学习了 Unity 客户端开发、Python asyncio 网络编程、TCP/UDP 通信、服务器权威状态、客户端预测和状态同步等内容。")
add_para(doc, "目前项目可以在本机或局域网环境中运行多人模式，基本功能已经能够完整演示。后续如果继续完善，可以加入断线重连、房间持久化、公网服务器部署、更加高效的二进制协议以及更完善的压力测试。")
add_para(doc, "以上就是我们项目的整体介绍，谢谢大家。")

add_heading(doc, "十一、建议的录制顺序")
for i, text in enumerate([
    "主菜单和项目名称",
    "单机模式设置怪物数量和血量",
    "移动、射击、怪物击退和死亡爆炸",
    "血条、死亡倒计时和随机复活",
    "弹药、自动换弹和长按 R 补充弹药",
    "三波怪物和倒计时提示",
    "小地图和击杀排行榜",
    "启动服务器，展示两个客户端同步",
    "说明 TCP、UDP 和客户端平滑处理",
    "以项目亮点和总结结束",
], 1):
    add_para(doc, f"{i}. {text}")

doc.core_properties.title = "Unity 多人合作射击项目视频演讲稿"
doc.core_properties.subject = "项目内容、功能演示和设计亮点介绍"
doc.core_properties.author = "项目组"
doc.save(OUT)
print(OUT)
